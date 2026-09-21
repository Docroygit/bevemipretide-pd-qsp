"""
Merge Narrative_Review.docx + Methods_and_Results_v2.docx -> Manuscript_new.docx
- Justified text alignment throughout
- AI watermark removal
- Section 1 (Narrative Review) → Table A → Section 2 (Methods) → Section 3 (Results)
  → Table B → Section 4 (Discussion) → References
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from copy import deepcopy
import os, re

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "New publication plots")

# AI watermark phrases to fix
AI_FIXES = [
    ("The broader landscape of PD disease-modification trials provides "
     "instructive context.",
     "Published PD disease-modification trials offer relevant comparisons."),
    ("leverage early biomarker signals",
     "use early biomarker signals"),
    ("actionable roadmap",
     "concrete basis"),
]


def get_para_text(element):
    """Extract text from a w:p XML element."""
    return ''.join(t.text for t in element.iter(qn('w:t')) if t.text)


def is_heading(element):
    """Check if a w:p element uses a Heading style."""
    pPr = element.find(qn('w:pPr'))
    if pPr is None:
        return False
    pStyle = pPr.find(qn('w:pStyle'))
    if pStyle is None:
        return False
    val = pStyle.get(qn('w:val')) or ''
    return 'Heading' in val


def merge():
    nr = Document(os.path.join(OUT_DIR, "Narrative_Review.docx"))
    mr = Document(os.path.join(OUT_DIR, "Methods_and_Results_v2.docx"))

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 2.0
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    body = doc.element.body
    # Remove default empty paragraph(s)
    for child in list(body):
        if child.tag.endswith('}p'):
            body.remove(child)

    sectPr = body.find(qn('w:sectPr'))

    def insert(el):
        clone = deepcopy(el)
        if sectPr is not None:
            sectPr.addprevious(clone)
        else:
            body.append(clone)

    # ── Part 1: Narrative review (before "References") ──────────────
    hit_refs = False
    refs_elements = []
    for el in nr.element.body:
        if el.tag.endswith('}sectPr'):
            continue
        if not hit_refs and el.tag.endswith('}p'):
            if is_heading(el) and 'References' in get_para_text(el):
                hit_refs = True
                refs_elements.append(el)
                continue
        if hit_refs:
            refs_elements.append(el)
            continue
        insert(el)

    # ── Part 2: Table A + Methods + Results + Table B + Discussion ──
    for el in mr.element.body:
        if el.tag.endswith('}sectPr'):
            continue
        insert(el)

    # ── Part 3: References (from narrative review) ──────────────────
    for el in refs_elements:
        insert(el)

    # ── Apply justified alignment to all body paragraphs ────────────
    for p in doc.paragraphs:
        if not p.style.name.startswith('Heading'):
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # ── Fix AI watermarks ───────────────────────────────────────────
    fix_count = 0
    for p in doc.paragraphs:
        for run in p.runs:
            for old, new in AI_FIXES:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    fix_count += 1

    # ── Save ────────────────────────────────────────────────────────
    path = os.path.join(OUT_DIR, "Manuscript_new.docx")
    doc.save(path)
    print(f"Saved: {path}")
    print(f"AI watermark fixes applied: {fix_count}")

    # ── Verify structure ────────────────────────────────────────────
    headings = []
    for p in doc.paragraphs:
        if p.style.name.startswith('Heading'):
            headings.append(p.text.strip())
    print("\nDocument structure (headings):")
    for h in headings:
        print(f"  {h}")

    # ── Word counts by section ──────────────────────────────────────
    current_sec = "Preamble"
    sec_wc = {}
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style.name.startswith('Heading'):
            current_sec = txt
            continue
        if current_sec == "References":
            continue
        wc = len(txt.split())
        sec_wc[current_sec] = sec_wc.get(current_sec, 0) + wc

    total = 0
    for sec, wc in sec_wc.items():
        print(f"  {sec}: {wc}w")
        total += wc
    print(f"  TOTAL (excl. references): {total}w")

    # ── Verify tables survived ──────────────────────────────────────
    print(f"\nTables in document: {len(doc.tables)}")
    for i, tbl in enumerate(doc.tables):
        rows = len(tbl.rows)
        cols = len(tbl.columns)
        first_cell = tbl.rows[0].cells[0].text[:40]
        print(f"  Table {i+1}: {rows}×{cols}, header: '{first_cell}...'")


if __name__ == "__main__":
    merge()
